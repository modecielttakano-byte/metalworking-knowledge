#!/usr/bin/env python3
"""Build a local SQLite FTS index from two metalworking Excel workbooks.

Expected relative paths (from repo root):

  data/internal.xlsx   社内ナレッジブック
  data/qa.xlsx         外部 Q&A（シート Q&Aデータ）
  data/knowledge.db    出力

Company workbooks and the DB are gitignored. Do not commit them.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DEFAULT_INTERNAL = ROOT / "data" / "internal.xlsx"
DEFAULT_QA = ROOT / "data" / "qa.xlsx"
DEFAULT_DB = ROOT / "data" / "knowledge.db"

# --- source labels (UI) -------------------------------------------------
SRC_LOG = "internal_log"
SRC_TS = "internal_ts_tool"
SRC_QA = "external_qa"

SOURCE_LABEL = {
    SRC_LOG: "社内ログ",
    SRC_TS: "社内TS・工具",
    SRC_QA: "外部Q&A",
}
TRUST_INTERNAL = "社内実績"
TRUST_FORUM = "掲示板意見"

LOG_SHEETS = {
    "SWG",
    "FPC",
    "COIL",
    "ARM",
    "ARR",
    "BH",
    "PVT",
    "OL1",
    "OL2",
    "SLIT",
    "TAP",
    "PH",
    "SH",
    "DT",
    "OD",
    "SWS",
    "Drill",
    "Chamfer",
    "APP",
    "JIG",
    "Coolant",
    "Loader",
}

HEADER_TO_FIELD = {
    "機種名": "model",
    "加工機械": "machine",
    "工場": "factory",
    "発生日": "occurred_on",
    "実施日": "acted_on",
    "現象、問題点": "phenomenon",
    "現象、問題点": "phenomenon",
    "現象・問題点": "phenomenon",
    "確認内容": "phenomenon",
    "対応検討": "discussion",
    "実施対策": "countermeasure",
    "結果": "result",
}

# Columns / cells that are author-owner style and must not be searchable.
PII_CELL_RE = re.compile(
    r"^(作成者|担当者|記入者|責任者|承認者|確認者)\s*[:：]",
)
PII_BY_RE = re.compile(r"^By\s+[A-Za-z][A-Za-z.\-]*\s*$")
PII_SAN_RE = re.compile(r"[\u4e00-\u9fff]{1,4}さん")
PII_INLINE_AUTHOR_RE = re.compile(r"(作成者|担当者)\s*[:：]\s*\S+")

QA_PROCESS_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("リーマ", ("リーマ", "リーマー", "reamer")),
    ("ドリル", ("ドリル", "drill")),
    ("エンドミル", ("エンドミル", "endmill", "end mill")),
    ("タップ", ("タップ", "tap ")),
    ("面取り", ("面取り", "面取", "チャンファ", "chamfer")),
    ("ワイヤ放電", ("ワイヤ", "ワイヤーカット", "wire cut")),
    ("放電加工", ("放電", "edm")),
    ("旋盤", ("旋盤", "lathe")),
    ("フライス", ("フライス", "マシニング")),
    ("研削", ("研削", "研磨")),
    ("クーラント", ("クーラント", "切削油")),
]

SCHEMA = """
PRAGMA journal_mode = WAL;
PRAGMA foreign_keys = ON;

DROP TABLE IF EXISTS documents_fts;
DROP TABLE IF EXISTS documents;
DROP TABLE IF EXISTS meta;

CREATE TABLE documents (
    id INTEGER PRIMARY KEY,
    source_type TEXT NOT NULL,
    source_label TEXT NOT NULL,
    trust_label TEXT NOT NULL,
    process TEXT,
    sheet_name TEXT,
    title TEXT NOT NULL,
    body TEXT NOT NULL,
    model TEXT,
    machine TEXT,
    factory TEXT,
    tool TEXT,
    url TEXT,
    has_result INTEGER NOT NULL DEFAULT 0,
    has_answer INTEGER NOT NULL DEFAULT 0,
    extra_json TEXT
);

CREATE VIRTUAL TABLE documents_fts USING fts5(
    ngrams,
    content='',
    tokenize='unicode61 remove_diacritics 2'
);

CREATE TABLE meta (
    key TEXT PRIMARY KEY,
    value TEXT
);

CREATE INDEX idx_docs_source ON documents(source_type);
CREATE INDEX idx_docs_process ON documents(process);
CREATE INDEX idx_docs_has_result ON documents(has_result);
CREATE INDEX idx_docs_has_answer ON documents(has_answer);
"""


def nfkc(text: str) -> str:
    return unicodedata.normalize("NFKC", text or "")


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.time):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return ""
        return value.strftime("%H:%M")
    if isinstance(value, dt.datetime):
        if value.year < 1901:
            return ""
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value == int(value):
            return str(int(value))
        return str(value)
    text = str(value).strip()
    if text in {"None", "00:00:00", "nan"}:
        return ""
    return nfkc(text)


def is_pii_cell(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if PII_CELL_RE.match(t) or PII_BY_RE.match(t):
        return True
    # Signature block labels with no technical content
    if t in {"承認", "確認", "作成", "APPROVE", "CHECK", "PREPARE"}:
        return True
    return False


def mask_pii(text: str) -> str:
    if not text:
        return ""
    text = PII_INLINE_AUTHOR_RE.sub(r"\1：（非表示）", text)
    text = re.sub(r"By\s+[A-Za-z][A-Za-z.\-]*", "By （非表示）", text)
    text = PII_SAN_RE.sub("（氏名）", text)
    return text


def clean_text(value: Any) -> str:
    text = mask_pii(cell_str(value))
    if is_pii_cell(text):
        return ""
    return text


def join_unique(parts: Iterable[str], sep: str = "\n") -> str:
    seen: set[str] = set()
    out: list[str] = []
    for part in parts:
        t = (part or "").strip()
        if not t or t in seen:
            continue
        seen.add(t)
        out.append(t)
    return sep.join(out)


def to_ngrams(text: str) -> str:
    """Character 1-grams + 2-grams for Japanese substring search via FTS5."""
    text = nfkc(mask_pii(text or ""))
    tokens: list[str] = []
    buf: list[str] = []

    def flush() -> None:
        if not buf:
            return
        piece = "".join(buf)
        buf.clear()
        if not piece:
            return
        ascii_word = all(ord(ch) < 128 for ch in piece)
        if ascii_word:
            tokens.append(piece.lower())
            if len(piece) >= 2:
                for i in range(len(piece) - 1):
                    tokens.append(piece[i : i + 2].lower())
            return
        for ch in piece:
            tokens.append(ch)
        if len(piece) >= 2:
            for i in range(len(piece) - 1):
                tokens.append(piece[i : i + 2])

    for ch in text:
        if ch.isalnum() or ch in "ーｰ":
            buf.append(ch)
        else:
            flush()
    flush()
    # FTS operators must not appear as tokens
    cleaned = []
    for tok in tokens:
        if tok.upper() in {"AND", "OR", "NOT", "NEAR"}:
            continue
        if any(c in tok for c in '"^:*()'):
            continue
        cleaned.append(tok)
    return " ".join(cleaned)


def infer_qa_process(title: str, body: str, answer: str) -> str:
    blob = f"{title}\n{body}\n{answer}".lower()
    blob_nfkc = nfkc(blob)
    for label, keys in QA_PROCESS_RULES:
        for key in keys:
            if key.lower() in blob_nfkc.lower():
                return label
    return ""


class IndexWriter:
    def __init__(self, db_path: Path) -> None:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        if db_path.exists():
            db_path.unlink()
        for extra in (str(db_path) + "-wal", str(db_path) + "-shm"):
            Path(extra).unlink(missing_ok=True)
        self.conn = sqlite3.connect(str(db_path))
        self.conn.executescript(SCHEMA)
        self.count = 0

    def add(
        self,
        *,
        source_type: str,
        process: str,
        sheet_name: str,
        title: str,
        body: str,
        model: str = "",
        machine: str = "",
        factory: str = "",
        tool: str = "",
        url: str = "",
        has_result: bool = False,
        has_answer: bool = False,
        extra: dict[str, Any] | None = None,
    ) -> None:
        title = mask_pii(nfkc(title or "")).strip()
        body = mask_pii(nfkc(body or "")).strip()
        if not title and not body:
            return
        if not title:
            title = body.split("\n", 1)[0][:80]
        searchable = " ".join(
            p
            for p in (
                title,
                body,
                process,
                model,
                machine,
                factory,
                tool,
                sheet_name,
            )
            if p
        )
        ngrams = to_ngrams(searchable)
        if not ngrams:
            return
        extra_json = json.dumps(extra or {}, ensure_ascii=False)
        cur = self.conn.execute(
            """
            INSERT INTO documents (
                source_type, source_label, trust_label, process, sheet_name,
                title, body, model, machine, factory, tool, url,
                has_result, has_answer, extra_json
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                source_type,
                SOURCE_LABEL[source_type],
                TRUST_INTERNAL if source_type != SRC_QA else TRUST_FORUM,
                process or "",
                sheet_name or "",
                title[:200],
                body,
                model or "",
                machine or "",
                factory or "",
                tool or "",
                url or "",
                1 if has_result else 0,
                1 if has_answer else 0,
                extra_json,
            ),
        )
        doc_id = cur.lastrowid
        self.conn.execute(
            "INSERT INTO documents_fts(rowid, ngrams) VALUES (?, ?)",
            (doc_id, ngrams),
        )
        self.count += 1

    def finish(self, stats: dict[str, Any]) -> None:
        self.conn.execute(
            "INSERT INTO meta(key, value) VALUES ('stats', ?)",
            (json.dumps(stats, ensure_ascii=False),),
        )
        self.conn.execute(
            "INSERT INTO meta(key, value) VALUES ('built_at', ?)",
            (dt.datetime.now().isoformat(timespec="seconds"),),
        )
        self.conn.commit()
        self.conn.close()


def iter_rows(ws, max_row: int | None = None):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if max_row and i > max_row:
            break
        yield i, row


def ingest_problem_logs_single_pass(ws, sheet_name: str, writer: IndexWriter) -> int:
    mapping: dict[int, str] = {}
    process = sheet_name
    added = 0
    current: dict[str, list[str]] | None = None
    seen_header = False
    skip_next_if_english = False

    def flush() -> None:
        nonlocal added, current
        if not current:
            return
        def cat(key: str) -> str:
            return join_unique(current.get(key, []))

        phen = cat("phenomenon")
        cm = cat("countermeasure")
        result = cat("result")
        disc = cat("discussion")
        if not (phen or cm or result):
            current = None
            return
        body_parts = []
        if phen:
            body_parts.append(f"【現象・問題点】\n{phen}")
        if disc:
            body_parts.append(f"【対応検討】\n{disc}")
        if cm:
            body_parts.append(f"【実施対策】\n{cm}")
        if result:
            body_parts.append(f"【結果】\n{result}")
        title = phen.split("\n")[0][:80] if phen else (cm.split("\n")[0][:80] if cm else process)
        writer.add(
            source_type=SRC_LOG,
            process=process,
            sheet_name=sheet_name,
            title=title,
            body="\n\n".join(body_parts),
            model=cat("model"),
            machine=cat("machine"),
            factory=cat("factory"),
            has_result=bool(result),
            extra={
                "occurred_on": cat("occurred_on"),
                "acted_on": cat("acted_on"),
                "has_countermeasure": bool(cm),
            },
        )
        added += 1
        current = None

    for rno, row in iter_rows(ws):
        vals = [clean_text(c) for c in row]
        nonempty = [v for v in vals if v]
        if not seen_header:
            if rno <= 4 and nonempty:
                titled = [v for v in nonempty if v.upper() != "TOP"]
                if titled:
                    process = titled[0]
            joined = " ".join(vals)
            if "機種名" in joined and (
                "実施対策" in joined or "結果" in joined or "確認内容" in joined
            ):
                seen_header = True
                for col, val in enumerate(vals):
                    field = HEADER_TO_FIELD.get(val)
                    if field:
                        mapping[col] = field
                skip_next_if_english = True
            continue
        if skip_next_if_english:
            skip_next_if_english = False
            joined = " ".join(vals).lower()
            if "model" in joined or "machine" in joined or "phenomenon" in joined:
                continue

        parsed: dict[str, str] = {}
        for col, field in mapping.items():
            if col < len(vals) and vals[col]:
                parsed[field] = vals[col]

        core = any(parsed.get(k) for k in ("phenomenon", "discussion", "countermeasure", "result"))
        if not core:
            flush()
            continue
        if current is None:
            current = {k: [v] for k, v in parsed.items() if v}
            continue
        # Blank rows already flushed. Remaining rows in a block are one incident
        # (phenomenon / countermeasure often wrap across Excel rows).
        for k, v in parsed.items():
            if not v:
                continue
            current.setdefault(k, [])
            if v not in current[k]:
                current[k].append(v)
    flush()
    return added


def ingest_troubleshooting(ws, sheet_name: str, tool_kind: str, writer: IndexWriter) -> int:
    added = 0
    last_phen = ""
    header_seen = False
    for rno, row in iter_rows(ws):
        vals = [clean_text(c) for c in row]
        if not header_seen:
            if "現象" in vals and "原因" in vals:
                header_seen = True
            continue
        # Typical layout: col2=index, col3=phenomenon, col5=cause, col7=action
        phen = vals[2] if len(vals) > 2 else ""
        cause = vals[4] if len(vals) > 4 else ""
        action = vals[6] if len(vals) > 6 else ""
        # Some rows put bullets in adjacent columns; pick longest-ish text fields
        if not cause:
            # fallback: any medium-length cell besides phen
            cands = [v for v in vals if v and v not in {phen, "●", "・"} and len(v) >= 2]
            if len(cands) >= 2:
                cause, action = cands[0], cands[1] if len(cands) > 1 else ""
            elif len(cands) == 1:
                cause = cands[0]
        if phen and phen not in {"●", "・"}:
            last_phen = phen
        if not (cause or action):
            continue
        if not last_phen:
            continue
        title = f"{tool_kind}：{last_phen}"
        body = f"【現象】\n{last_phen}\n\n【原因】\n{cause}\n\n【対策】\n{action}"
        writer.add(
            source_type=SRC_TS,
            process=tool_kind,
            sheet_name=sheet_name,
            title=title,
            body=body,
            tool=tool_kind,
            extra={"cause": cause, "action": action},
        )
        added += 1
    return added


def ingest_endmill(ws, writer: IndexWriter) -> int:
    mapping: dict[int, str] = {}
    header_seen = False
    skip_en = False
    groups: list[dict[str, list[str]]] = []
    current: dict[str, list[str]] | None = None
    field_of = {
        "メーカー": "maker",
        "型番": "model_no",
        "径": "diameter",
        "刃長": "flute_length",
        "刃数": "flutes",
        "角度": "angle",
        "コート": "coating",
        "選考理由": "reason",
        "加工内容": "process_desc",
        "粗さ（直線）": "ra_line",
        "粗さ（R)": "ra_r",
        "粗さ（R）": "ra_r",
        "結果": "result",
    }

    def flush() -> None:
        nonlocal current
        if current:
            groups.append(current)
            current = None

    for rno, row in iter_rows(ws):
        vals = [clean_text(c) for c in row]
        if not header_seen:
            if "メーカー" in vals and "型番" in vals:
                header_seen = True
                for col, val in enumerate(vals):
                    if val in field_of:
                        mapping[col] = field_of[val]
                skip_en = True
            continue
        if skip_en:
            skip_en = False
            if "Maker" in vals or "maker" in " ".join(vals).lower():
                continue
        parsed = {}
        for col, field in mapping.items():
            if col < len(vals) and vals[col]:
                parsed[field] = vals[col]
        maker = parsed.get("maker", "")
        model_no = parsed.get("model_no", "")
        if not parsed:
            continue
        if maker or model_no:
            key_changed = (
                current is not None
                and (
                    (maker and maker not in current.get("maker", []))
                    or (model_no and model_no not in current.get("model_no", []))
                )
            )
            if current is None or key_changed:
                flush()
                current = {}
        if current is None:
            current = {}
        for k, v in parsed.items():
            current.setdefault(k, [])
            if v not in current[k]:
                current[k].append(v)
    flush()

    added = 0
    for g in groups:
        def cat(key: str, sep: str = " / ") -> str:
            return join_unique(g.get(key, []), sep=sep)

        maker = cat("maker")
        model_no = cat("model_no")
        if not (maker or model_no):
            continue
        result = cat("result", "\n")
        reason = cat("reason", "\n")
        process_desc = cat("process_desc", "\n")
        tool = " / ".join(p for p in (maker, model_no) if p)
        title = f"エンドミル選定：{tool}"
        specs = " / ".join(
            p
            for p in (
                cat("diameter") and f"径 {cat('diameter')}",
                cat("flute_length") and f"刃長 {cat('flute_length')}",
                cat("flutes") and f"刃数 {cat('flutes')}",
                cat("angle") and f"角度 {cat('angle')}",
                cat("coating") and f"コート {cat('coating')}",
            )
            if p
        )
        parts = [f"【工具】\n{tool}"]
        if specs:
            parts.append(f"【仕様】\n{specs}")
        if reason:
            parts.append(f"【選考理由】\n{reason}")
        if process_desc:
            parts.append(f"【加工内容】\n{process_desc}")
        if result:
            parts.append(f"【結果】\n{result}")
        writer.add(
            source_type=SRC_TS,
            process="エンドミル",
            sheet_name="mil",
            title=title,
            body="\n\n".join(parts),
            tool=tool,
            has_result=bool(result),
            extra={"maker": maker, "model_no": model_no},
        )
        added += 1
    return added


def ingest_text_note(
    ws,
    sheet_name: str,
    writer: IndexWriter,
    *,
    process: str,
    title: str,
    skip_values: set[str] | None = None,
) -> int:
    skip_values = skip_values or {"TOP", "戻る", "Back"}
    chunks: list[str] = []
    seen: set[str] = set()
    for rno, row in iter_rows(ws, max_row=80):
        for cell in row:
            text = clean_text(cell)
            if not text or text in skip_values or text in seen:
                continue
            if is_pii_cell(text):
                continue
            if len(text) < 2:
                continue
            if text.replace(".", "", 1).replace("-", "", 1).isdigit():
                continue
            seen.add(text)
            chunks.append(text)
    if not chunks:
        return 0
    body = "\n".join(chunks)
    writer.add(
        source_type=SRC_TS,
        process=process,
        sheet_name=sheet_name,
        title=title,
        body=body,
    )
    return 1


def ingest_reamer_conditions(ws, writer: IndexWriter) -> int:
    lines: list[str] = []
    for rno, row in iter_rows(ws, max_row=20):
        vals = [clean_text(c) for c in row[:18]]
        if not any(vals):
            continue
        lines.append(" ".join(v for v in vals if v))
    if not lines:
        return 0
    writer.add(
        source_type=SRC_TS,
        process="リーマ",
        sheet_name="リーマ加工条件計算",
        title="リーマ推奨加工条件（アルミニウム）",
        body="\n".join(lines),
        tool="リーマ",
    )
    return 1


def ingest_internal(path: Path, writer: IndexWriter) -> dict[str, int]:
    stats: dict[str, int] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        log_n = 0
        for name in wb.sheetnames:
            if name in LOG_SHEETS:
                n = ingest_problem_logs_single_pass(wb[name], name, writer)
                log_n += n
        stats["internal_log"] = log_n

        ts_n = 0
        if "ドリルトラブルシューティング" in wb.sheetnames:
            ts_n += ingest_troubleshooting(
                wb["ドリルトラブルシューティング"],
                "ドリルトラブルシューティング",
                "ドリル",
                writer,
            )
        if "リーマトラブルシューティング" in wb.sheetnames:
            ts_n += ingest_troubleshooting(
                wb["リーマトラブルシューティング"],
                "リーマトラブルシューティング",
                "リーマ",
                writer,
            )
        if "mil" in wb.sheetnames:
            ts_n += ingest_endmill(wb["mil"], writer)
        if "リーマ加工条件計算" in wb.sheetnames:
            ts_n += ingest_reamer_conditions(wb["リーマ加工条件計算"], writer)
        if "Root & Tip Step" in wb.sheetnames:
            ts_n += ingest_text_note(
                wb["Root & Tip Step"],
                "Root & Tip Step",
                writer,
                process="Root & Tip",
                title="Root & Tip / Slit カッターマーク確認",
            )
        if "穴径大" in wb.sheetnames:
            ts_n += ingest_text_note(
                wb["穴径大"],
                "穴径大",
                writer,
                process="リーマ",
                title="リーマ交換時の対処（穴径大）",
            )
        if "穴挽き目" in wb.sheetnames:
            ts_n += ingest_text_note(
                wb["穴挽き目"],
                "穴挽き目",
                writer,
                process="リーマ",
                title="リーマ交換時の対処（穴挽き目）",
            )
        if "Step" in wb.sheetnames:
            ts_n += ingest_text_note(
                wb["Step"],
                "Step",
                writer,
                process="スリット／チップステップ",
                title="スリット・チップステップ対策メモ",
            )
        if "Mil2" in wb.sheetnames:
            ts_n += ingest_text_note(
                wb["Mil2"],
                "Mil2",
                writer,
                process="エンドミル",
                title="厚物ACB向けエンドミル改善トライ",
            )
        # Picture captions only (no binary media)
        for pic_sheet in ("Picture", "Pic"):
            if pic_sheet in wb.sheetnames:
                ts_n += ingest_text_note(
                    wb[pic_sheet],
                    pic_sheet,
                    writer,
                    process="外観・刃具写真メモ",
                    title=f"{pic_sheet}シートのキャプション",
                    skip_values={"TOP", "戻る", "Back"},
                )
        stats["internal_ts_tool"] = ts_n
    finally:
        wb.close()
    return stats


def ingest_qa(path: Path, writer: IndexWriter) -> dict[str, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    added = 0
    answered = 0
    try:
        sheet = wb[wb.sheetnames[0]]
        header = None
        for rno, row in iter_rows(sheet):
            vals = [cell_str(c) for c in row]
            if header is None:
                header = [v.strip() for v in vals]
                continue
            rec = {header[i]: vals[i] if i < len(vals) else "" for i in range(len(header))}
            url = rec.get("URL", "")
            title = rec.get("質問タイトル", "") or "無題"
            question = rec.get("質問本文", "")
            answer = rec.get("回答", "")
            category = rec.get("カテゴリ", "")
            if not (title or question or answer):
                continue
            has_answer = bool(answer.strip())
            process = infer_qa_process(title, question, answer)
            body_parts = [f"【質問】\n{question}"]
            if has_answer:
                body_parts.append(f"【回答】\n{answer}")
            else:
                body_parts.append("【回答】\n（未回答）")
            writer.add(
                source_type=SRC_QA,
                process=process or category or "機械加工",
                sheet_name="Q&Aデータ",
                title=title[:120],
                body="\n\n".join(body_parts),
                url=url,
                has_answer=has_answer,
                extra={"category": category},
            )
            added += 1
            if has_answer:
                answered += 1
    finally:
        wb.close()
    return {"external_qa": added, "external_qa_answered": answered}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="加工方法ナレッジの Excel を SQLite に取り込む")
    parser.add_argument(
        "--internal",
        type=Path,
        default=DEFAULT_INTERNAL,
        help="社内ブックのパス（相対パス例: data/internal.xlsx）",
    )
    parser.add_argument(
        "--qa",
        type=Path,
        default=DEFAULT_QA,
        help="外部Q&Aブックのパス（相対パス例: data/qa.xlsx）",
    )
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="出力 SQLite")
    parser.add_argument(
        "--sample",
        action="store_true",
        help="sample_data/ の合成ブックを取り込む（デモ用）",
    )
    args = parser.parse_args(argv)

    if args.sample:
        args.internal = ROOT / "sample_data" / "internal_sample.xlsx"
        args.qa = ROOT / "sample_data" / "qa_sample.xlsx"

    missing = []
    if not args.internal.exists():
        missing.append(str(args.internal))
    if not args.qa.exists():
        missing.append(str(args.qa))
    if missing:
        print("Excel が見つかりません:", file=sys.stderr)
        for m in missing:
            print(f"  - {m}", file=sys.stderr)
        print(
            "\nリポジトリ直下で次のように配置してください:\n"
            "  data/internal.xlsx  … 社内ナレッジブック\n"
            "  data/qa.xlsx        … 外部 Q&A（シート名 Q&Aデータ）\n"
            "デモだけ試す場合: python ingest.py --sample",
            file=sys.stderr,
        )
        return 1

    print(f"internal: {args.internal}")
    print(f"qa:       {args.qa}")
    print(f"db:       {args.db}")
    writer = IndexWriter(args.db)
    stats: dict[str, Any] = {}
    stats.update(ingest_internal(args.internal, writer))
    stats.update(ingest_qa(args.qa, writer))
    stats["total"] = writer.count
    writer.finish(stats)
    print("取り込み完了:")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
