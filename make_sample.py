#!/usr/bin/env python3
"""Create tiny synthetic workbooks for the public demo (not real shop data)."""

from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "sample_data"


def write_internal(path: Path) -> None:
    wb = Workbook()
    # TOP
    top = wb.active
    top.title = "TOP"
    top["B2"] = "サンプル目次（実データではありません）"

    swg = wb.create_sheet("SWG")
    swg["G2"] = "SWG-H"
    swg["J2"] = "TOP"
    headers = ["", "機種名", "加工機械", "工場", "発生日", "実施日", "現象、問題点", "対応検討", "実施対策", "結果"]
    for col, h in enumerate(headers, 1):
        swg.cell(5, col, h)
    en = ["", "Model", "Machine", "Factory", "Occurrence", "Actual", "Phenomenon", "", "Countermeasure", "Result"]
    for col, h in enumerate(en, 1):
        swg.cell(6, col, h)
    swg["B7"] = "SAMPLE-A"
    swg["C7"] = "CNC"
    swg["D7"] = "DemoPlant"
    swg["G7"] = "穴抜け側にバリが残る"
    swg["I7"] = "貫通直前の送りを落とす。面取り工程を追加。"
    swg["J7"] = "バリ高さは管理値内に入った"
    swg["B8"] = "SAMPLE-A"
    swg["C8"] = "CNC"
    swg["D8"] = "DemoPlant"
    swg["I8"] = "リーマの送りを F200→F80 に変更"

    drill_log = wb.create_sheet("Drill")
    drill_log["G2"] = "Tool Broken"
    for col, h in enumerate(["", "機種名", "加工機械", "工場", "発生日", "実施日", "確認内容", "実施対策", "結果"], 1):
        drill_log.cell(5, col, h)
    drill_log["B7"] = "SAMPLE-B"
    drill_log["C7"] = "IMM"
    drill_log["D7"] = "DemoPlant"
    drill_log["G7"] = "下穴ドリル折損"
    drill_log["H7"] = "ステップ量を小さくし、回転を下げる"
    drill_log["I7"] = "折損は再発せず"

    mil = wb.create_sheet("mil")
    mil["F2"] = "Endmil"
    mil_h = ["", "メーカー", "型番", "径", "刃長", "刃数", "角度", "コート", "選考理由", "加工内容", "粗さ（直線）", "", "粗さ（R)", "", "結果"]
    for col, h in enumerate(mil_h, 1):
        mil.cell(5, col, h)
    mil["B7"] = "デモメーカ"
    mil["C7"] = "DEMO-EM5"
    mil["D7"] = "φ5"
    mil["E7"] = 25
    mil["F7"] = 4
    mil["I7"] = "送りを上げてもひきめが残らないかを確認"
    mil["J7"] = "外周仕上げ"
    mil["O7"] = "送り F300 ではひきめNG。F150 で可"

    ts = wb.create_sheet("ドリルトラブルシューティング")
    ts["B2"] = "トラブルシューティング（ドリル）"
    ts["C4"] = "現象"
    ts["E4"] = "原因"
    ts["G4"] = "対策"
    ts["C5"] = "穴抜けバリが大きい"
    ts["E5"] = "貫通時の送りが大きい"
    ts["G5"] = "貫通時、送りを下げる"
    ts["C6"] = "ドリルが折損する"
    ts["E6"] = "切りくず詰まり"
    ts["G6"] = "ステップフィードの回数を多くする"

    rts = wb.create_sheet("リーマトラブルシューティング")
    rts["B2"] = "トラブルシューティング（リーマ）"
    rts["C4"] = "現象"
    rts["E4"] = "原因"
    rts["G4"] = "対策"
    rts["C5"] = "穴が大きくなった場合"
    rts["E5"] = "切削速度または送り速度が高すぎる"
    rts["G5"] = "切削速度または送り速度を下げる"
    rts["E6"] = "面取り（チャンファー）の振れが大きい"
    rts["G6"] = "チャンファーの再研磨またはツール交換"

    cond = wb.create_sheet("リーマ加工条件計算")
    cond["B2"] = "サンプル推奨加工条件"
    cond["B4"] = "被削材"
    cond["C4"] = "アルミニウム"
    cond["B7"] = "周速"
    cond["C7"] = 25

    # Author cell must be stripped by ingest (PII mask check)
    mil2 = wb.create_sheet("Mil2")
    mil2["I1"] = "作成者：山田　太郎"
    mil2["C3"] = "エンドミル加工改善トライ（サンプル）"
    mil2["H3"] = "目的 ： サンプル厚物の外周品質"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_qa(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Q&Aデータ"
    ws.append(["URL", "カテゴリ", "質問タイトル", "質問本文", "回答", "ベストアンサー"])
    ws.append(
        [
            "https://example.invalid/qa/reamer-burr",
            "機械加工",
            "リーマ加工でバリが出る",
            "アルミにリーマを通すと穴抜け側にバリが残ります。送りと回転の目安は？",
            "貫通際の送りを落とす、面取りを先に入れる、切れ味の落ちたリーマを使わない、がよくある打ち手です。",
            None,
        ]
    )
    ws.append(
        [
            "https://example.invalid/qa/chamfer",
            "機械加工",
            "面取り量の決め方",
            "ドリル穴の面取りをC0.2にするかC0.5にするか迷っています。",
            "組立干渉とバリ取り目的で決まります。バリ取りだけなら小さい面取り＋手直しでも足りることがあります。",
            None,
        ]
    )
    ws.append(
        [
            "https://example.invalid/qa/feed-unanswered",
            "機械加工",
            "送り速度の上限がわからない",
            "φ6エンドミルでアルミを削るときの送り上限を知りたいです。",
            None,
            None,
        ]
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    write_internal(OUT / "internal_sample.xlsx")
    write_qa(OUT / "qa_sample.xlsx")
    print(f"wrote {OUT / 'internal_sample.xlsx'}")
    print(f"wrote {OUT / 'qa_sample.xlsx'}")
